"""v3.1 第1项：基于折外预测的相对评分、预测支持度分档、250款主表与12张评分卡。

只读取已有的5次重复OOF预测与逐维评价，不重新训练模型。
规则见 docs/解题思路.md 第1.4–1.7节。
AI assistance: Claude (Claude Opus 5.5), Anthropic, 2026-10-26. 用于代码实现，详见支撑材料《AI 工具使用详情说明》。
"""
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
FEAS = ROOT / 'analysis' / 'A_feasibility'
OUT = ROOT / 'results' / 'q1'
OUT.mkdir(parents=True, exist_ok=True)

# ---------- 读取已有结果 ----------
oof = np.load(FEAS / 'floor_ablation' / 'baseline_oof.npz', allow_pickle=True)
pred, beers, fields = oof['prediction'], oof['beer'].astype(int), list(oof['fields'])
assert pred.shape == (5, len(beers), len(fields))
splits = json.load(open(FEAS / 'floor_ablation' / 'split_and_floor_log.json', encoding='utf-8'))
assert len(splits) == 15

y = pd.read_csv(ROOT / '4.csv').set_index('beer')
assert list(y.columns[2:]) == fields
y = y[fields]
meta = pd.read_csv(FEAS / 'day1' / 'sensory_metadata.csv', encoding='utf-8-sig').set_index('R_name')
types = meta.loc[fields, 'DataType']
numeric = [f for f in fields if types[f] == 'numeric']
categorical = [f for f in fields if types[f] == 'categorical']
dims = pd.read_csv(FEAS / 'q1_metrics' / 'ExtraTrees_dimension_summary.csv', encoding='utf-8-sig', index_col=0).loc[fields]


def parent_of(col):
    return '_'.join(col.split('_')[:2]) + '_all'


for c in categorical:
    assert types[parent_of(c)] == 'numeric'

# ---------- 参考预测：每款啤酒5次重复OOF均值 ----------
ref = pd.DataFrame(pred.mean(axis=0), index=beers, columns=fields)
ref.index.name = 'beer'

# 每款啤酒在其作为测试集的5个划分中，对应训练折的真实均值，再跨重复平均
train_mean = pd.DataFrame(0.0, index=beers, columns=fields)
count = pd.Series(0, index=beers)
test_seen = {r: set() for r in range(5)}
for s in splits:
    tr = s['train_beer']
    te = sorted(set(beers) - set(tr))
    test_seen[s['repeat']].update(te)
    train_mean.loc[te] += y.loc[tr].mean().values
    count.loc[te] += 1
assert (count == 5).all() and all(len(v) == len(beers) for v in test_seen.values())
train_mean = train_mean.div(count, axis=0)

# ---------- numeric：P5/P95 锚点（冻结） ----------
q05, q95 = ref[numeric].quantile(0.05), ref[numeric].quantile(0.95)
no_info = (q95 - q05) <= 1e-12
anchors = pd.DataFrame({'q05_oof': q05, 'q95_oof': q95, 'no_discrimination': no_info})


def relative_score(z):
    """z: DataFrame of standardized predictions for numeric fields -> integer 0–6 (NaN if no discrimination)."""
    s = 6 * (z[numeric] - q05) / (q95 - q05)
    s = s.clip(0, 6).round()
    s.loc[:, no_info[no_info].index] = np.nan
    return s


score_num = relative_score(ref)

# ---------- categorical：相对勾选倾向 + 父项门控 ----------
tick = pd.DataFrame(index=ref.index, columns=categorical, dtype=int)
for c in categorical:
    tick[c] = ((score_num[parent_of(c)] > 0) & (ref[c] > train_mean[c])).astype(int)
display = pd.DataFrame(index=ref.index, columns=categorical, dtype=float)
for c in categorical:
    display[c] = tick[c] * score_num[parent_of(c)]

# ---------- 预测支持度分档 ----------
r2 = dims['R2']
tier = pd.Series(np.where(r2 >= 0.3, 'A', np.where(r2 >= 0.1, 'B', 'C')), index=fields)
support = pd.DataFrame({
    'type': types, 'R2_mean_oof': r2, 'Spearman': dims['Spearman'], 'MAE': dims['MAE'],
    'MSE_skill_vs_style': dims['MSE_skill_vs_style'],
    'support_tier': tier,
    'skill_positive': dims['MSE_skill_vs_style'] > 0,
})
support['flag'] = np.where(~support.skill_positive, '未优于风格均值', '')
support.loc[categorical, 'final_tick_rate'] = tick.mean()
support.loc[categorical, 'parent_positive_rate'] = [(score_num[parent_of(c)] > 0).mean() for c in categorical]
support.loc[numeric, 'score_dist'] = [
    json.dumps(score_num[c].value_counts().sort_index().astype(int).to_dict()) for c in numeric]
support = support.join(anchors)

# ---------- 写出主表 ----------
style = pd.read_csv(ROOT / '1.csv').set_index('beer')['tasting_category_fine']
ref.join(style).to_csv(OUT / 'q1_oof_mean_standardized_250.csv', encoding='utf-8-sig')
train_mean.to_csv(OUT / 'q1_train_fold_mean_reference_250.csv', encoding='utf-8-sig')
score_num.join(style).to_csv(OUT / 'q1_numeric_relative_scores_250.csv', encoding='utf-8-sig')
tick.join(style).to_csv(OUT / 'q1_categorical_ticks_250.csv', encoding='utf-8-sig')
display.join(style).to_csv(OUT / 'q1_categorical_template_display_250.csv', encoding='utf-8-sig')
support.to_csv(OUT / 'q1_prediction_support_tiers.csv', encoding='utf-8-sig')
anchors.to_json(OUT / 'q1_frozen_anchors.json', force_ascii=False, indent=1, orient='index')

# ---------- 填写12张评分卡（不改动原始3.xlsx） ----------
# 强度行（各自的0–6标记行号）
A_ROWS = {'malt': 7, 'hops': 20, 'esters': 33}
SUB_ROWS = {  # 子项勾选行，顺序与 4.csv 字段一致
    'malt': {'grain': 10, 'bread': 12, 'cara': 14, 'burn': 16},
    'hops': {'citrus': 23, 'tropical': 25, 'noble': 27, 'woody': 29},
    'esters': {'ethac': 36, 'isoaa': 38, 'flower': 40, 'fruity': 42},
}
LEFT_GRADE_COLS = ['C', 'E', 'G', 'I', 'K', 'M', 'O']
RIGHT_GRADE_COLS = ['S', 'U', 'W', 'Y', 'AA', 'AC', 'AE']
SINGLE = {  # field -> (row, side)
    'body': (50, 'L'), 'alcohol': (50, 'R'), 'bitternes': (54, 'L'), 'sweetness': (54, 'R'),
    'co2': (58, 'L'), 'acidity': (58, 'R'), 'aftertaste': (62, 'L'),
    'dms': (68, 'L'), 't2n': (68, 'R'), 'metallic': (72, 'L'), 'diacetyl': (72, 'R'),
    'X4vg': (76, 'L'), 'stale_hops': (76, 'R'),
    'orange': (82, 'L'), 'lactic': (82, 'R'), 'coriander': (86, 'L'), 'acetic': (86, 'R'),
    'clove': (90, 'L'), 'barnyard': (90, 'R'), 'overall': (94, 'L'),
}
MARK = 'X'
wb = openpyxl.load_workbook(ROOT / '3.xlsx')
filled = []
for k in range(1, 13):
    ws = wb[f'beer_{k}']
    b = int(ws['S3'].value)
    assert b == k and b in score_num.index
    grade_cols = {'L': LEFT_GRADE_COLS, 'R': RIGHT_GRADE_COLS}

    def put_grade(row, side, val):
        if pd.isna(val):
            return
        ws[f'{grade_cols[side][int(val)]}{row}'] = MARK

    for side, prefix in (('L', 'A'), ('R', 'F')):
        for grp, row in A_ROWS.items():
            put_grade(row, side, score_num.loc[b, f'{prefix}_{grp}_all'])
            for sub, srow in SUB_ROWS[grp].items():
                if tick.loc[b, f'{prefix}_{grp}_{sub}'] == 1:
                    ws[f'{"C" if side == "L" else "S"}{srow}'] = MARK
    for f, (row, side) in SINGLE.items():
        put_grade(row, side, score_num.loc[b, f])
    c_fields = [f for f in fields if tier[f] == 'C']
    ws['S95'] = ('相对评分（基于折外预测，非原始品评分）。预测支持度C档（参考分）：' + ', '.join(c_fields))
    filled.append(b)

# 模板 summary 表中 barnyard 列（AX）误引用了 Stale hops 单元格 $AG$76，应为 $AG$90
fixes = []
sm = wb['summary']
for r in range(2, 14):
    cell = sm[f'AX{r}']
    if isinstance(cell.value, str) and cell.value.endswith('!$AG$76'):
        cell.value = cell.value.replace('!$AG$76', '!$AG$90')
        fixes.append(cell.coordinate)
wb.save(OUT / '3_filled_v31.xlsx')

# 评分卡对应数值（便于不经 Excel 重算直接核对）
cards = score_num.loc[filled].join(display.loc[filled]).join(style)
cards.to_csv(OUT / 'q1_scorecards_12_values.csv', encoding='utf-8-sig')

summary = {
    'rule_numeric': 'round(clip(6*(z-q05_oof)/(q95_oof-q05_oof),0,6)); anchors from 250 five-repeat OOF means, frozen',
    'rule_categorical': 'tick = 1(parent relative score>0) * 1(OOF mean > train-fold mean averaged over the beer\'s 5 test folds)',
    'tier_rule': 'A: R2>=0.3, B: 0.1<=R2<0.3, C: R2<0.1 (5-repeat mean OOF R2)',
    'tier_counts': tier.value_counts().sort_index().to_dict(),
    'tier_counts_numeric': tier[numeric].value_counts().sort_index().to_dict(),
    'tier_counts_categorical': tier[categorical].value_counts().sort_index().to_dict(),
    'skill_nonpositive_fields': support.index[~support.skill_positive].tolist(),
    'no_discrimination_fields': no_info[no_info].index.tolist(),
    'numeric_score_distribution': score_num.stack().value_counts(normalize=True).sort_index().round(4).to_dict(),
    'final_tick_rate_range': [float(tick.mean().min()), float(tick.mean().max())],
    'final_tick_rate_median': float(tick.mean().median()),
    'scorecards_filled_beer': filled,
    'template_formula_fixes': fixes,
    'template_rows_left_blank': ['Acetaldehyde (row 45): no corresponding target in 4.csv'],
}
json.dump(summary, open(OUT / 'summary.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
print(json.dumps(summary, ensure_ascii=False, indent=1, default=str))
